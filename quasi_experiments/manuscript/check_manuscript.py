"""Validate the generated submission package.

The target journal is selected by the JOURNAL environment variable
(default: policy_sciences). Journal rules live in JOURNAL_CONFIG.
"""

from __future__ import annotations

import csv
import hashlib
import importlib
import json
import math
import os
import re
import sys
import zipfile
from pathlib import Path

import pandas as pd
import requests
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
QROOT = ROOT / "quasi_experiments"
MANUSCRIPT = QROOT / "manuscript"
JOURNAL = os.environ.get("JOURNAL", "policy_sciences")

JOURNAL_CONFIG = {
    "policy_sciences": {
        "name": "Policy Sciences",
        "abstract_range": (150, 250),
        "keywords_range": (4, 6),
        "word_limit": 10000,
        "title_marker": "Article type: Analysis",
        "expected_main_tables": ["1", "2", "3", "4"],
        "guideline_note": (
            "Guideline target: 150-250 word abstract, 4-6 keywords, no more than "
            "three heading levels, APA-style author-date references, Word source, "
            "automatic page numbers, and a 10,000-word maximum."
        ),
        "article_note": "Policy Sciences article type selected: Analysis.",
    },
    "policy_studies": {
        "name": "Policy Studies",
        "abstract_range": (150, 200),
        "keywords_range": (5, 6),
        "word_limit": 9000,
        "title_marker": "Journal: Policy Studies",
        "expected_main_tables": ["1", "2", "3", "4", "5"],
        "guideline_note": (
            "Guideline target: approximately 150-200 word abstract, 5-6 keywords, "
            "no more than three heading levels, APA-style author-date references, "
            "Word source, automatic page numbers, and approximately 9,000 words "
            "all-inclusive."
        ),
        "article_note": "Policy Studies (Taylor & Francis) original research article.",
    },
    "jcpa": {
        "name": "Journal of Comparative Policy Analysis",
        "abstract_range": (150, 250),
        "keywords_range": (5, 6),
        "word_limit": 8000,
        "title_marker": "Journal: Journal of Comparative Policy Analysis",
        "expected_main_tables": ["1", "2", "3", "4", "5"],
        "guideline_note": (
            "Guideline target: explicit comparative analysis, abstract under "
            "250 words, 5-6 keywords, no more than three heading levels, "
            "APA-style author-date references, Word source, double-anonymized "
            "manuscript, and an 8,000-word maximum."
        ),
        "article_note": (
            "JCPA (Taylor & Francis) original research article; explicitly "
            "comparative design required."
        ),
    },
}
CFG = JOURNAL_CONFIG[JOURNAL]

MAIN_DOCX = ROOT / f"manuscript_{JOURNAL}.docx"
SEPARATE_DOCX = ROOT / f"manuscript_{JOURNAL}_figures_separate.docx"
APPENDIX_DOCX = ROOT / f"appendix_{JOURNAL}.docx"
TITLE_DOCX = ROOT / f"title_page_{JOURNAL}.docx"
COVER_DOCX = ROOT / f"cover_letter_{JOURNAL}.docx"
STATEMENTS_DOCX = ROOT / f"submission_statements_{JOURNAL}.docx"
REPORT = MANUSCRIPT / f"check_report_{JOURNAL}.md"
sys.path.insert(0, str(MANUSCRIPT))
textmod = importlib.import_module(f"text_{JOURNAL}")


def docx_text(path: Path) -> tuple[str, str]:
    doc = Document(path)
    ordinary = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            ordinary.extend(cell.text for cell in row.cells)
    with zipfile.ZipFile(path) as archive:
        xml = "\n".join(
            archive.read(name).decode("utf-8", errors="replace")
            for name in archive.namelist()
            if name.endswith(".xml")
        )
    return "\n".join(ordinary), xml


def document_word_count(path: Path) -> int:
    text, _ = docx_text(path)
    return len(re.findall(r"\b[\w'-]+\b", text))


def nonascii_scan(paths: list[Path]) -> list[str]:
    hits = []
    for path in paths:
        if path.suffix == ".docx":
            source, _ = docx_text(path)
        elif path.suffix == ".pptx":
            with zipfile.ZipFile(path) as archive:
                source = "\n".join(
                    "".join(
                        re.findall(
                            r"<a:t[^>]*>(.*?)</a:t>",
                            archive.read(name).decode("utf-8", errors="replace"),
                        )
                    )
                    for name in archive.namelist()
                    if name.startswith("ppt/slides/") and name.endswith(".xml")
                )
        else:
            source = path.read_text(encoding="utf-8", errors="replace")
        for line_no, line in enumerate(source.splitlines(), start=1):
            chars = [(char, ord(char)) for char in line if ord(char) > 127]
            if chars:
                hits.append(f"{path}:{line_no}: {chars}")
    return hits


def cited_objects(text: str) -> tuple[list[str], list[str]]:
    figures = list(dict.fromkeys(re.findall(r"\bFigure\s+([A-Z]?\d+)", text)))
    tables = list(dict.fromkeys(re.findall(r"\bTable\s+([A-Z]?\d+)", text)))
    return figures, tables


def extract_abstract_and_keywords(path: Path) -> tuple[str, list[str]]:
    paragraphs = [paragraph.text.strip() for paragraph in Document(path).paragraphs]
    abstract_index = paragraphs.index("Abstract")
    keywords_index = next(
        index for index, value in enumerate(paragraphs) if value.startswith("Keywords:")
    )
    abstract = " ".join(paragraphs[abstract_index + 1 : keywords_index])
    keywords = [
        value.strip()
        for value in paragraphs[keywords_index].split("Keywords:", 1)[1].split(";")
        if value.strip()
    ]
    return abstract, keywords


def reference_citation_notes(body: str, numbers: dict[str, object]) -> list[str]:
    notes = []
    aliases = {
        "International Commission on Radiological Protection": "ICRP",
        "United Nations Scientific Committee on the Effects of Atomic Radiation": "UNSCEAR",
        "World Health Organization": "WHO",
    }
    for reference in (
        value.format(**numbers) for value in textmod.REFERENCES
    ):
        author = reference.split(" (", 1)[0].split(",", 1)[0].rstrip(".")
        year_match = re.search(r"\((\d{4})", reference)
        if not year_match:
            notes.append(f"Reference year could not be parsed: {reference[:80]}")
            continue
        year = year_match.group(1)
        token = aliases.get(author, author)
        token_pattern = re.escape(token).replace(r"\ ", r"\s+")
        if not re.search(
            rf"(?:{token_pattern}.{{0,120}}{year}|{year}.{{0,120}}{token_pattern})",
            body,
            flags=re.I | re.S,
        ):
            notes.append(f"Reference may be uncited: {author} ({year})")
    return notes


def verify_reference_links(numbers: dict[str, object]) -> list[str]:
    notes = []
    for reference in (
        value.format(**numbers) for value in textmod.REFERENCES
    ):
        doi = re.search(r"10\.\d{4,9}/\S+", reference)
        if doi:
            clean_doi = doi.group(0).rstrip(".")
            try:
                response = requests.get(
                    f"https://api.crossref.org/works/{clean_doi}", timeout=15
                )
                notes.append(f"Crossref verification {response.status_code}: {clean_doi}")
            except requests.RequestException as exc:
                notes.append(f"Crossref unverifiable: {clean_doi} ({exc})")
            continue
        urls = re.findall(r"https?://\S+", reference)
        if not urls:
            continue
        url = urls[0].rstrip(".")
        try:
            response = requests.get(url, timeout=15)
            notes.append(f"URL verification {response.status_code}: {url}")
        except requests.RequestException as exc:
            notes.append(f"URL unverifiable: {url} ({exc})")
    return notes


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def validate_raw_manifest(
    project: str, failures: list[str], notes: list[str]
) -> None:
    data_dir = QROOT / project / "data"
    raw_dir = data_dir / "raw"
    manifest = data_dir / "raw_manifest.csv"
    if not manifest.exists():
        failures.append(f"Missing raw-data checksum manifest: {project}.")
        return
    with manifest.open(encoding="utf-8", newline="") as stream:
        rows = list(csv.DictReader(stream))
    raw_files = sorted(path for path in raw_dir.iterdir() if path.is_file())
    row_paths = {row["path"] for row in rows}
    expected_paths = {str(path.relative_to(QROOT)) for path in raw_files}
    if row_paths != expected_paths:
        missing = sorted(expected_paths - row_paths)
        extra = sorted(row_paths - expected_paths)
        failures.append(
            f"{project} raw manifest coverage mismatch; missing={missing}, extra={extra}."
        )
        return
    for row in rows:
        path = QROOT / row["path"]
        if int(row["bytes"]) != path.stat().st_size:
            failures.append(f"Raw-data size mismatch: {row['path']}.")
        if row["sha256"] != sha256(path):
            failures.append(f"Raw-data SHA-256 mismatch: {row['path']}.")
        if row["source_url"] == "UNMATCHED":
            failures.append(f"Raw-data source is unmatched: {row['path']}.")
        for field in (
            "acquisition_conditions",
            "usage_terms",
            "timestamp_precision",
        ):
            if not row[field]:
                failures.append(f"Raw-data manifest field {field} is blank: {row['path']}.")
    legacy_unknown = sum(
        row["source_url"] == "UNKNOWN_LEGACY_UNUSED" for row in rows
    )
    notes.append(
        f"{project} raw-data manifest verified: {len(rows)} files, "
        f"{legacy_unknown} unused legacy file(s) with unavailable source metadata."
    )


def validate_cross_domain_manifest(
    failures: list[str], notes: list[str]
) -> None:
    data_dir = ROOT / "data"
    manifest = data_dir / "source_manifest.csv"
    if not manifest.exists():
        failures.append("Missing cross-domain source-data checksum manifest.")
        return
    with manifest.open(encoding="utf-8", newline="") as stream:
        rows = list(csv.DictReader(stream))
    source_files = sorted(
        path
        for path in data_dir.glob("*.csv")
        if path.name not in {"source_manifest.csv", "source_metadata.csv"}
    )
    row_paths = {row["path"] for row in rows}
    expected_paths = {str(path.relative_to(ROOT)) for path in source_files}
    if row_paths != expected_paths:
        missing = sorted(expected_paths - row_paths)
        extra = sorted(row_paths - expected_paths)
        failures.append(
            "Cross-domain source manifest coverage mismatch; "
            f"missing={missing}, extra={extra}."
        )
        return
    for row in rows:
        path = ROOT / row["path"]
        if int(row["bytes"]) != path.stat().st_size:
            failures.append(f"Source-data size mismatch: {row['path']}.")
        if row["sha256"] != sha256(path):
            failures.append(f"Source-data SHA-256 mismatch: {row['path']}.")
        for field in (
            "source_url",
            "data_identifier",
            "acquisition_timestamp_utc",
            "timestamp_precision",
            "acquisition_conditions",
            "usage_terms",
        ):
            if not row[field]:
                failures.append(
                    f"Cross-domain source manifest field {field} is blank: "
                    f"{row['path']}."
                )
    notes.append(
        f"Cross-domain source-data manifest verified: {len(rows)} files."
    )


def main() -> int:
    failures: list[str] = []
    notes: list[str] = []
    numbers = json.loads((MANUSCRIPT / "numbers.json").read_text())

    for project in ("kansas", "fukushima"):
        validate_raw_manifest(project, failures, notes)
    validate_cross_domain_manifest(failures, notes)

    expected_files = [
        MAIN_DOCX,
        SEPARATE_DOCX,
        APPENDIX_DOCX,
        TITLE_DOCX,
        ROOT / f"tables_{JOURNAL}.docx",
        ROOT / f"figures_{JOURNAL}.pptx",
        COVER_DOCX,
        STATEMENTS_DOCX,
    ]
    for path in expected_files:
        if not path.exists():
            failures.append(f"Missing generated submission file: {path.name}")
    if failures:
        REPORT.write_text(
            "# Manuscript check report\n\n## Failures\n"
            + "\n".join(f"- {failure}" for failure in failures)
            + "\n",
            encoding="utf-8",
        )
        print(REPORT.read_text())
        return 1

    text, xml = docx_text(MAIN_DOCX)
    body = text.split("References", 1)[0]
    for marker in ["bougtoir", "Onishi", "Shiga University", "gmail.com"]:
        if marker.lower() in text.lower():
            failures.append(f"Anonymization keyword found in main DOCX: {marker}")
    if numbers["repo_url"] in text:
        failures.append("Public author-identifying repository URL found in main DOCX.")

    abstract, keywords = extract_abstract_and_keywords(MAIN_DOCX)
    abstract_words = len(re.findall(r"\b[\w'-]+\b", abstract))
    lo, hi = CFG["abstract_range"]
    if not lo <= abstract_words <= hi:
        failures.append(
            f"Abstract has {abstract_words} words; {CFG['name']} target is {lo}-{hi}."
        )
    klo, khi = CFG["keywords_range"]
    if not klo <= len(keywords) <= khi:
        failures.append(
            f"Keyword count is {len(keywords)}; {CFG['name']} requires {klo}-{khi}."
        )

    all_docx = [path for path in expected_files if path.suffix == ".docx"]
    source_files = [
        MANUSCRIPT / "build_manuscript.py",
        MANUSCRIPT / f"text_{JOURNAL}.py",
        MANUSCRIPT / "check_manuscript.py",
    ]
    hits = nonascii_scan(
        all_docx + [ROOT / f"figures_{JOURNAL}.pptx"] + source_files
    )
    if hits:
        failures.append("Non-ASCII text found outside OMML:")
        failures.extend(f"- {hit}" for hit in hits)

    figures, tables = cited_objects(body)
    expected_main_figures = [str(index) for index in range(1, 8)]
    expected_appendix_figures = ["A1"]
    expected_main_tables = CFG["expected_main_tables"]
    expected_appendix_tables = ["A1", "A2"]
    main_figures = [value for value in figures if not value.startswith("A")]
    appendix_figures = [value for value in figures if value.startswith("A")]
    main_tables = [value for value in tables if not value.startswith("A")]
    appendix_tables = [value for value in tables if value.startswith("A")]
    if (
        main_figures != expected_main_figures
        or appendix_figures != expected_appendix_figures
    ):
        failures.append(
            "Figure first-mention order is "
            f"main={main_figures}, appendix={appendix_figures}; expected "
            f"main={expected_main_figures}, appendix={expected_appendix_figures}."
        )
    if main_tables != expected_main_tables or appendix_tables != expected_appendix_tables:
        failures.append(
            "Table first-mention order is "
            f"main={main_tables}, appendix={appendix_tables}; expected "
            f"main={expected_main_tables}, appendix={expected_appendix_tables}."
        )

    pngs = sorted((MANUSCRIPT / "figures").glob("Figure*.png"))
    tiffs = sorted((MANUSCRIPT / "figures").glob("Figure*.tiff"))
    if len(pngs) != 8:
        failures.append(f"Expected eight manuscript PNGs, found {len(pngs)}.")
    if len(tiffs) != 8:
        failures.append(f"Expected eight manuscript TIFFs, found {len(tiffs)}.")

    doc = Document(MAIN_DOCX)
    if JOURNAL == "jcpa":
        paragraphs = [paragraph.text.strip() for paragraph in doc.paragraphs]
        if not any(p.startswith("Appendix") for p in paragraphs):
            failures.append(
                "Appendix section missing from the main manuscript DOCX."
            )
        if len(doc.inline_shapes) != 8:
            failures.append(
                "Expected eight inline figures (7 main + Appendix A1) in "
                f"the main manuscript, found {len(doc.inline_shapes)}."
            )
        if len(doc.tables) != len(expected_main_tables) + 2:
            failures.append(
                f"Expected {len(expected_main_tables) + 2} inline tables "
                "(5 main + Appendix A1/A2) in the main manuscript, found "
                f"{len(doc.tables)}."
            )
    else:
        if "Appendix" in [paragraph.text.strip() for paragraph in doc.paragraphs]:
            failures.append("Appendix heading found in the main manuscript DOCX.")
        if len(doc.inline_shapes) != 7:
            failures.append(
                f"Expected seven inline figures in the main manuscript, found {len(doc.inline_shapes)}."
            )
        if len(doc.tables) != len(expected_main_tables):
            failures.append(
                f"Expected {len(expected_main_tables)} inline tables in the main "
                f"manuscript, found {len(doc.tables)}."
            )
    appendix = Document(APPENDIX_DOCX)
    if len(appendix.inline_shapes) != 1:
        failures.append(
            f"Expected one inline appendix figure, found {len(appendix.inline_shapes)}."
        )
    if len(appendix.tables) != 2:
        failures.append(
            f"Expected two appendix tables, found {len(appendix.tables)}."
        )
    for generated in [doc, appendix]:
        heading_levels = [
            int(match.group(1))
            for paragraph in generated.paragraphs
            if (match := re.match(r"Heading (\d+)", paragraph.style.name))
        ]
        if heading_levels and max(heading_levels) > 3:
            failures.append("A generated DOCX uses more than three heading levels.")

    if "m:oMath" not in xml:
        failures.append("No native Word OMML equation found in the main DOCX.")
    if " PAGE " not in xml:
        failures.append("Automatic page-number field not found in the main DOCX.")
    if "Artificial intelligence assistance" not in text:
        failures.append("AI-use disclosure heading is missing from the main DOCX.")
    if textmod.AI_ASSISTANCE_TEXT not in text:
        failures.append("The synchronized AI-use declaration is missing from the main DOCX.")
    if "Data availability" not in text:
        failures.append("Data availability statement is missing from the main DOCX.")

    total_word_count = document_word_count(MAIN_DOCX)
    if total_word_count > CFG["word_limit"]:
        failures.append(
            f"Manuscript word count is {total_word_count}; {CFG['name']} limit is "
            f"{CFG['word_limit']:,}."
        )
    title_text, _ = docx_text(TITLE_DOCX)
    title_count_match = re.search(
        r"Manuscript word count \(all inclusive\): ([\d,]+)", title_text
    )
    if not title_count_match:
        failures.append("Title page does not report the all-inclusive word count.")
    elif int(title_count_match.group(1).replace(",", "")) != total_word_count:
        failures.append(
            "Title-page word count does not match the generated main manuscript."
        )
    if CFG["title_marker"] not in title_text:
        failures.append(
            f"Title page is missing the expected marker: {CFG['title_marker']}"
        )
    separate_text, _ = docx_text(SEPARATE_DOCX)
    statements_text, _ = docx_text(STATEMENTS_DOCX)
    cover_text, _ = docx_text(COVER_DOCX)
    for label, generated_text in [
        ("separate-figure manuscript", separate_text),
        ("title page", title_text),
        ("submission statements", statements_text),
    ]:
        if textmod.AI_ASSISTANCE_TEXT not in generated_text:
            failures.append(
                f"The synchronized AI-use declaration is missing from the {label}."
            )
    for label, generated_text in [
        ("main manuscript", text),
        ("separate-figure manuscript", separate_text),
        ("cover letter", cover_text),
    ]:
        if re.search(r"\bsix[- ](?:step|stage|layer)s?\b", generated_text, re.I):
            failures.append(f"Ambiguous framework counting remains in the {label}.")

    for key, value in numbers.items():
        if key == "manuscript_word_count":
            continue
        value_text = str(value)
        if value_text not in text and value_text.replace("< ", "<") not in text:
            notes.append(
                f"numbers.json value not found verbatim in main DOCX: {key}={value_text}"
            )

    summary = pd.read_csv(ROOT / "results" / "summary_table.csv")
    preferred = pd.read_csv(
        QROOT / "fukushima/results/did_permutation_inference.csv"
    ).query(
        "treatment == 'T1' and control_pool == 'fukushima_only' "
        "and weighting == 'unweighted' and age_control == False"
    ).iloc[0]
    independently_derived = {
        "share_not_robust_pct": int(
            round(100 * (summary["Verdict"] != "ROBUST_NONLINEAR").mean())
        ),
        "t1_main_att": preferred["att"],
        "t1_main_perm_p": preferred["perm_p_att"],
        "t1_excess_2011": pd.read_csv(
            QROOT / "fukushima/results/acute_2011_decomposition.csv"
        )
        .query("group == 'T1' and muni_name == 'pooled'")
        .iloc[0]["excess_2011"],
        "ks_iit_pct": int(
            round(
                100
                * (
                    1
                    - math.exp(
                        pd.read_csv(QROOT / "kansas/results/sc_summary.csv")
                        .query("outcome == 'ln_real_iit_pc' and pool == 'main'")
                        .iloc[0]["post_mean_gap"]
                    )
                )
            )
        ),
        "ks_iit_perm_p": pd.read_csv(
            QROOT / "kansas/results/sc_summary.csv"
        )
        .query("outcome == 'ln_real_iit_pc' and pool == 'main'")
        .iloc[0]["perm_p"],
    }
    for key, value in independently_derived.items():
        if key == "t1_excess_2011":
            expected = f"{value:,.0f}"
        elif key == "t1_main_att":
            expected = f"{value:.3f}"
        elif key == "t1_main_perm_p":
            expected = "< 0.001" if value < 0.001 else f"{value:.3f}"
        elif key == "ks_iit_perm_p":
            expected = f"{value:.2f}"
        else:
            expected = str(value)
        if str(numbers[key]) != expected:
            failures.append(
                f"Independent derivation mismatch for {key}: "
                f"numbers={numbers[key]} derived={expected}."
            )

    notes.extend(reference_citation_notes(body, numbers))
    notes.extend(verify_reference_links(numbers))
    notes.extend(
        [
            CFG["article_note"],
            CFG["guideline_note"],
            "The main manuscript is anonymized; author details and the public repository URL are confined to editorial submission files.",
            "Measured municipal mortality, certified disaster-related deaths, and projected radiological deaths remain distinct estimands.",
        ]
    )

    report = [
        f"# {CFG['name']} manuscript check report",
        "",
        f"- Abstract word count: {abstract_words}",
        f"- Keyword count: {len(keywords)}",
        f"- All-inclusive main manuscript word count: {total_word_count}",
        f"- Main inline figures: {len(doc.inline_shapes)}",
        f"- Main inline tables: {len(doc.tables)}",
        "",
        "## Independent headline values",
    ]
    report.extend(f"- {key}: {value}" for key, value in independently_derived.items())
    report.extend(["", "## Notes"])
    report.extend(f"- {note}" for note in notes)
    report.extend(["", "## Failures"])
    report.extend(f"- {failure}" for failure in failures)
    if not failures:
        report.append("- None.")
    REPORT.write_text("\n".join(report) + "\n", encoding="utf-8")
    print(REPORT.read_text())
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
