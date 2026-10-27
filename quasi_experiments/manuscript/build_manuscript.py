"""Build the journal manuscript and submission files.

The target journal is selected by the JOURNAL environment variable
(default: policy_sciences). Journal-specific manuscript text lives in
manuscript/text_<journal>.py.
"""

from __future__ import annotations

import importlib
import json
import math
import os
import re
import sys
from pathlib import Path
from typing import Any

import matplotlib.pyplot as plt
import matplotlib.patches as patches
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches as PptInches, Pt as PptPt

ROOT = Path(__file__).resolve().parents[2]
QROOT = ROOT / "quasi_experiments"
MANUSCRIPT = QROOT / "manuscript"
FIGURES = MANUSCRIPT / "figures"
RESULTS = ROOT / "results"
FUKU_RESULTS = QROOT / "fukushima" / "results"
FUKU_DATA = QROOT / "fukushima" / "data"
FUKU_RAW = FUKU_DATA / "raw"
KANSAS_RESULTS = QROOT / "kansas" / "results"
KANSAS_DATA = QROOT / "kansas" / "data"
sys.path.insert(0, str(MANUSCRIPT))
JOURNAL = os.environ.get("JOURNAL", "policy_sciences")
text = importlib.import_module(f"text_{JOURNAL}")

JOURNAL_META = {
    "policy_sciences": {"name": "Policy Sciences", "article_type": "Analysis"},
    "policy_studies": {
        "name": "Policy Studies",
        "article_type": "Original research article",
    },
    "jcpa": {
        "name": "Journal of Comparative Policy Analysis",
        "article_type": "Original research article",
    },
}

# Ordered table keys; filtered to keys present in both the generated frames
# and the journal's captions.
TABLE_KEYS = [
    "tab_audit",
    "tab_fuku_att",
    "tab_lnt",
    "tab_kansas",
    "tab_synthesis",
    "tab_a1",
    "tab_a2",
]

plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9})
PALETTE = ["#0072B2", "#D55E00", "#009E73", "#CC79A7", "#E69F00", "#56B4E9"]
ROMAJI_BY_CODE = {
    "07212": "Minamisoma",
    "07546": "Futaba",
    "07545": "Okuma",
    "07543": "Tomioka",
    "07308": "Kawamata",
    "07544": "Kawauchi",
    "07542": "Naraha",
    "07547": "Namie",
    "07548": "Katsurao",
    "07564": "Iitate",
    "07541": "Hirono",
    "07209": "Tamura",
}


def ptext(value: float, decimals: int = 3) -> str:
    if pd.isna(value):
        return "NA"
    if value == 0:
        return "< 0.001"
    if abs(value) < 0.001:
        return "< 0.001"
    return f"{value:.{decimals}f}"


def comma(value: float, decimals: int = 0) -> str:
    if decimals:
        return f"{value:,.{decimals}f}"
    return f"{int(round(value)):,}"


def signed_pct(value: float, decimals: int = 1) -> str:
    return f"{value:+.{decimals}f}"


def ascii_clean(value: Any) -> str:
    return str(value).replace("\u2013", "-").replace("\u2014", "-").replace("\u2212", "-")


def load_numbers() -> dict[str, Any]:
    summary = pd.read_csv(RESULTS / "summary_table.csv")
    sensitivity = pd.read_csv(RESULTS / "verdict_sensitivity_summary.csv")
    sample = json.loads((FUKU_RESULTS / "sample_summary.json").read_text())
    perm = pd.read_csv(FUKU_RESULTS / "did_permutation_inference.csv")
    att = pd.read_csv(FUKU_RESULTS / "att_summary.csv")
    event = pd.read_csv(FUKU_RESULTS / "did_event_study.csv")
    descriptive = pd.read_csv(FUKU_RESULTS / "descriptive_cdr.csv")
    acute = pd.read_csv(FUKU_RESULTS / "acute_2011_decomposition.csv")
    sc = pd.read_csv(FUKU_RESULTS / "sc_aggregate.csv").iloc[0]
    crossfit = pd.read_csv(FUKU_RESULTS / "sc_crossfit.csv")
    fuku_panel = pd.read_csv(
        FUKU_DATA / "municipal_panel.csv", dtype={"pref_code": str, "muni_code": str}
    )
    related = pd.read_csv(FUKU_DATA / "reconstruction_related_deaths_fukushima.csv")
    lnt = pd.read_csv(RESULTS.parent / "data" / "lnt_real.csv")
    ks = pd.read_csv(KANSAS_RESULTS / "sc_summary.csv")
    placebo = pd.read_csv(KANSAS_RESULTS / "placebo_in_time.csv")
    loo = pd.read_csv(KANSAS_RESULTS / "loo_robustness.csv")
    revenue = pd.read_csv(KANSAS_RESULTS / "revenue_counterfactual.csv")
    ks_meta = json.loads((KANSAS_RESULTS / "analysis_summary.json").read_text())
    fuku_codes = [str(code).zfill(5) for code in sample["t1_codes"]]
    fuku_names = (
        fuku_panel.loc[fuku_panel["muni_code"].astype(str).str.zfill(5).isin(fuku_codes), "muni_name"]
        .drop_duplicates()
        .tolist()
    )

    inventory_text = (ROOT / "README.md").read_text(encoding="utf-8")
    if "original 52 curves" not in inventory_text or "Yerkes-Dodson" not in inventory_text or "Weber-Fechner" not in inventory_text:
        raise RuntimeError("Inventory verification failed: README does not document 52 and 2.")
    n_curves = len(summary)
    counts = summary["Verdict"].value_counts().to_dict()
    preferred = perm.query(
        "treatment == 'T1' and control_pool == 'fukushima_only' "
        "and weighting == 'unweighted' and age_control == False"
    ).iloc[0]
    t1w_preferred = perm.query(
        "treatment == 'T1w' and control_pool == 'fukushima_only' "
        "and weighting == 'unweighted' and age_control == False"
    ).iloc[0]
    t2_preferred = perm.query(
        "treatment == 'T2' and control_pool == 'fukushima_only' "
        "and weighting == 'unweighted' and age_control == False"
    ).iloc[0]
    t1_event = event.query(
        "treatment == 'T1' and control_pool == 'fukushima_only' "
        "and weighting == 'unweighted' and age_control == False "
        "and event_year == 2011"
    ).iloc[0]
    t1_observed_post_deaths = fuku_panel.query(
        "year >= 2012 and year <= 2019 and muni_code in @fuku_codes"
    )["deaths"].sum()
    t1_rows = perm.query("treatment == 'T1'")
    t1_scope = att.query("treatment == 'T1' and adjustment == 'scope_robustness'").iloc[0]
    t1_acute = acute.query("group == 'T1' and muni_name == 'pooled'").iloc[0]
    crossfit_row = crossfit.query("treated == 'T1_aggregate'").iloc[0]
    t1_jump = descriptive.query("group == 'T1'").set_index("year")
    total_t1_pop = fuku_panel.query(
        "year == 2010 and muni_code in @fuku_codes"
    )["pop"].sum()
    ra_url = "https://www.reconstruction.go.jp/files/user/topics/main-cat2/sub-cat2-6/20260213_kanrenshi.pdf"
    fdr_url = "https://www.pref.fukushima.lg.jp/uploaded/attachment/422314.pdf"
    t1_excess_post = float(
        t1_observed_post_deaths * (1 - math.exp(-float(preferred["att"])))
    )

    main_iit = ks.query("outcome == 'ln_real_iit_pc' and pool == 'main'").iloc[0]
    all_iit = ks.query("outcome == 'ln_real_iit_pc' and pool == 'all'").iloc[0]
    main_tax = ks.query("outcome == 'ln_real_tax_pc' and pool == 'main'").iloc[0]
    main_gdp = ks.query("outcome == 'ln_rgdp_pc' and pool == 'main'").iloc[0]
    main_emp = ks.query("outcome == 'ln_emp' and pool == 'main'").iloc[0]
    placebo_macro = placebo.query("outcome in ['ln_rgdp_pc', 'ln_emp']")
    placebo_iit = placebo.query("outcome == 'ln_real_iit_pc'").iloc[0]
    iit_revenue = revenue.query("variable == 'real_iit_pc'")
    iit_loo = loo.query("outcome == 'ln_real_iit_pc'")["post_mean_gap"]
    lnt_row = summary.query("Curve == 'Linear No-Threshold (LNT)'").iloc[0]
    laffer_row = summary.query("Curve == 'Laffer Curve'").iloc[0]
    thj = pd.read_csv(FUKU_RESULTS / "lnt_comparison.csv")
    thj_row = thj.query("quantity.str.contains('ten Hoeve', na=False)", engine="python").iloc[0]
    ra_row = thj.query("quantity.str.contains('Reconstruction Agency', na=False)").iloc[0]
    thj_value = str(thj_row["value"])
    thj_range_match = re.search(r"([\d,]+)\s*[\u2013-]\s*([\d,]+)", thj_value)
    thj_best_match = re.search(r"best estimate\s+([\d,]+)", thj_value)
    if not thj_range_match or not thj_best_match:
        raise ValueError("Could not parse ten Hoeve and Jacobson estimate from lnt_comparison.csv")
    thj_range = (
        f"{int(thj_range_match.group(1).replace(',', '')):,}-"
        f"{int(thj_range_match.group(2).replace(',', '')):,}"
    )
    thj_best = int(thj_best_match.group(1).replace(",", ""))
    avoided_value = str(
        thj.query(
            "quantity.str.startswith('ten Hoeve and Jacobson evacuation-avoided')",
            engine="python",
        ).iloc[0]["value"]
    )
    avoided_range_match = re.search(r"([\d,]+)\s*[\u2013-]\s*([\d,]+)", avoided_value)
    avoided_best_match = re.search(r"best estimate\s+([\d,]+)", avoided_value)
    if not avoided_range_match or not avoided_best_match:
        raise ValueError(
            "Could not parse ten Hoeve evacuation-avoided estimate from "
            "lnt_comparison.csv"
        )
    thj_averted_range = (
        f"{int(avoided_range_match.group(1).replace(',', '')):,}-"
        f"{int(avoided_range_match.group(2).replace(',', '')):,}"
    )
    thj_averted_best = int(avoided_best_match.group(1).replace(",", ""))
    damage_text = (FUKU_RAW / "fukushima_damage_report_1260.txt").read_text(encoding="utf-8")
    fdr_date_match = re.search(
        r"\u5e73\u6210(\d+)\u5e74(\d+)\u6708(\d+)\u65e5", damage_text
    )
    fdr_year = 1988 + int(fdr_date_match.group(1))
    ra_date_match = re.search(r"/(\d{8})_", str(ra_row["source_url"]))
    ra_date = pd.to_datetime(ra_date_match.group(1), format="%Y%m%d").strftime("%B %-d, %Y")

    numbers = {
        "n_inventory": 52,
        "n_excluded": 2,
        "n_curves": n_curves,
        "n_not_robust": int((summary["Verdict"] != "ROBUST_NONLINEAR").sum()),
        "n_not_sig": int(counts.get("NOT_SIGNIFICANT", 0)),
        "n_outlier": int(counts.get("OUTLIER_DEPENDENT", 0)),
        "n_overfit": int(counts.get("OVERFITTING", 0)),
        "n_insufficient": int(counts.get("INSUFFICIENT_DATA", 0)),
        "share_not_robust_pct": int(round(100 * (summary["Verdict"] != "ROBUST_NONLINEAR").mean())),
        "share_min_pct": int(round(100 * sensitivity["share_not_robust"].min())),
        "share_max_pct": int(round(100 * sensitivity["share_not_robust"].max())),
        "alpha_default": 0.05,
        "k_default": 3,
        "alpha_grid": "{0.01, 0.05, 0.10}",
        "k_grid": "{1, ..., 5}",
        "laffer_n": int(laffer_row["N"]),
        "laffer_p": f"{laffer_row['p (full)']:.2f}",
        "lnt_n": int(lnt_row["N"]),
        "lnt_p": f"{lnt_row['p (full)']:.2f}",
        "lnt_n_lowdose": int((lnt["dose_gy"] < 0.1).sum()),
        "n_all_units": int(sample["n_all_units"]),
        "n_t1": len(fuku_names),
        "n_t1w": len(sample["treatment_groups"]["T1w"]),
        "n_t2": len(sample["treatment_groups"]["T2"]),
        "tsunami_cut": 1.20,
        "n_controls": int(sample["n_control_units"]),
        "n_dropped_missing": int(sample["n_dropped_missing"]),
        "n_dropped_top_prefs": (
            lambda d: "the drops are concentrated in "
            + " and ".join(f"{row.pref_name} ({int(row.n)})" for row in d.head(2).itertuples())
        )(
            pd.read_csv(FUKU_RESULTS / "dropped_units.csv")
            .groupby("pref_name")
            .size()
            .reset_index(name="n")
            .sort_values(["n", "pref_name"], ascending=[False, True])
        ),
        "evac_pop_approx": comma(round(total_t1_pop, -4)),
        "t1_main_att": f"{preferred['att']:.3f}",
        "t1_main_att_pct": int(round(100 * (1 - math.exp(float(preferred["att"]))))),
        "t1_main_perm_p": ptext(preferred["perm_p_att"]),
        "t1_main_pretrend_p": ptext(preferred["perm_p_pretrend"]),
        "t1_main_ev2011": f"{t1_event['beta']:.3f}",
        "t1_att_min": f"{t1_rows['att'].min():.3f}",
        "t1_att_max": f"{t1_rows['att'].max():.3f}",
        "t1_pct_min": int(round(100 * (1 - math.exp(float(t1_rows["att"].max()))))),
        "t1w_main_att": f"{t1w_preferred['att']:.3f}",
        "t1w_main_perm_p": ptext(t1w_preferred["perm_p_att"]),
        "t1w_pct_max": int(round(100 * (1 - math.exp(float(perm.query("treatment == 'T1w'")["att"].min()))))),
        "t2_perm_p_min": ptext(perm.query("treatment == 'T2'")["perm_p_att"].min()),
        "t1_pretrend_min": ptext(t1_rows["perm_p_pretrend"].min()),
        "t1_scope_att": f"{t1_scope['beta']:.3f}",
        "sc_agg_gap": f"{sc['post_mean_gap']:.3f}",
        "sc_agg_perm_p": ptext(sc["perm_p"]),
        "sc_agg_crossfit": f"{crossfit_row['evaluation_rmspe']:.3f}",
        "n_sc_donors": int(sc["n_donors"]),
        "sc_age_band": "10 percentage points",
        "n_sc_neg": int((pd.read_csv(FUKU_RESULTS / "sc_unit_summary.csv")["post_mean_gap"] < 0).sum()),
        "n_perm": 1000,
        "t1_2011_jump_pct": int(round(100 * (t1_jump.loc[2011, "cdr"] / t1_jump.loc[2010, "cdr"] - 1))),
        "t1_obs_2011": comma(t1_acute["deaths_2011"]),
        "t1_exp_2011": comma(t1_acute["expected_2011"]),
        "t1_excess_2011": comma(t1_acute["excess_2011"]),
        "t1_direct_2011": comma(t1_acute["direct_deaths_2011"]),
        "t1_resid_2011": comma(t1_acute["residual_2011"]),
        "t1_direct_share_pct": int(round(100 * t1_acute["direct_deaths_2011"] / t1_acute["excess_2011"])),
        "t1_main_excess_post": comma(abs(t1_excess_post)),
        "t1_main_excess_per_year": comma(abs(t1_excess_post) / 8),
        "ra_related_deaths": int(ra_row["value"]),
        "ra_date": ra_date,
        "ra_year": int(ra_date_match.group(1)[:4]),
        "ra_url": str(ra_row["source_url"]),
        "fdr_year": fdr_year,
        "fdr_no": 1260,
        "fdr_url": fdr_url,
        "thj_best": thj_best,
        "thj_range": thj_range,
        "thj_averted_best": thj_averted_best,
        "thj_averted_range": thj_averted_range,
        "ks_pre_years": "1997-2012",
        "ks_n_excluded": len(ks_meta["pools"]["rate_change_exclusions"]) + len(ks_meta["pools"]["no_broad_income_tax_exclusions"]),
        "ks_n_donors": int(len(ks_meta["pools"]["main"])),
        "ks_iit_pct": int(round(100 * (1 - math.exp(main_iit["post_mean_gap"])))),
        "ks_iit_perm_p": f"{main_iit['perm_p']:.2f}",
        "ks_iit_pct_all": int(round(100 * (1 - math.exp(all_iit["post_mean_gap"])))),
        "ks_iit_gap_2013": comma(abs(iit_revenue.query("year == 2013").iloc[0]["gap_per_capita_2012_dollars"])),
        "ks_iit_gap_2016": comma(abs(iit_revenue.query("year == 2016").iloc[0]["gap_per_capita_2012_dollars"])),
        "ks_iit_cum_bn": f"{abs(iit_revenue['counterfactual_gap_dollars'].sum()) / 1e9:.1f}",
        "ks_tax_pct": f"{abs(100 * (1 - math.exp(main_tax['post_mean_gap']))):.1f}",
        "ks_tax_perm_p": f"{main_tax['perm_p']:.2f}",
        "ks_gdp_pct": signed_pct(100 * (math.exp(main_gdp["post_mean_gap"]) - 1)),
        "ks_gdp_perm_p": f"{main_gdp['perm_p']:.2f}",
        "ks_emp_pct": signed_pct(100 * (math.exp(main_emp["post_mean_gap"]) - 1)),
        "ks_emp_perm_p": f"{main_emp['perm_p']:.2f}",
        "ks_placebo_min": f"{placebo_macro['post_mean_gap'].abs().min():.3f}",
        "ks_placebo_max": f"{placebo_macro['post_mean_gap'].abs().max():.3f}",
        "ks_iit_placebo": f"{abs(placebo_iit['post_mean_gap']):.3f}",
        "ks_iit_loo_min": f"{iit_loo.min():.3f}",
        "ks_iit_loo_max": f"{iit_loo.max():.3f}",
        "ks_iit_rmspe": f"{main_iit['rmspe_pre']:.3f}",
        "kansas_laffer_cite": "Rothschild, 2012",
        "repo_url": "https://github.com/bougtoir/canonical-curves-reexamination",
    }
    (MANUSCRIPT / "numbers.json").write_text(json.dumps(numbers, indent=2), encoding="utf-8")
    return numbers


def format_text(text: str, numbers: dict[str, Any]) -> str:
    return text.format(**numbers)


def set_style_font(doc: Document, style_name: str, size: float, bold: bool = False) -> None:
    style = doc.styles[style_name]
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.bold = bold
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    for style_name, size, bold in [
        ("Normal", 10, False),
        ("Title", 14, True),
        ("Heading 1", 12, True),
        ("Heading 2", 11, True),
        ("Caption", 9, False),
    ]:
        set_style_font(doc, style_name, size, bold)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Caption"].paragraph_format.space_before = Pt(12)
    doc.styles["Caption"].paragraph_format.space_after = Pt(6)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def document_word_count(path: Path) -> int:
    doc = Document(path)
    parts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return len(re.findall(r"\b[\w'-]+\b", "\n".join(parts)))


def save_fig(fig: plt.Figure, stem: str) -> None:
    FIGURES.mkdir(parents=True, exist_ok=True)
    fig.savefig(FIGURES / f"{stem}.png", dpi=300, bbox_inches="tight")
    fig.savefig(FIGURES / f"{stem}.tiff", dpi=300, bbox_inches="tight", pil_kwargs={"compression": "tiff_lzw"})
    plt.close(fig)


def plot_figures(numbers: dict[str, Any]) -> dict[str, Path]:
    paths: dict[str, Path] = {}
    FIGURES.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(7.0, 7.5))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    decisions = [
        "Nested F-test for the quadratic term:\np < alpha?",
        "Still p < alpha after removing the k most\ninfluential points (Cook's distance)?",
        "Leave-one-out CV:\nquadratic RMSE < linear RMSE?",
        "BIC prefers the quadratic model?",
    ]
    verdicts = ["Not significant", "Outlier dependent", "Overfitting", "BIC prefers linear"]
    ys = [0.86, 0.66, 0.46, 0.26]
    dx0, dw, dh = 0.06, 0.52, 0.12
    dcx = dx0 + dw / 2
    vx0, vw, vh = 0.72, 0.25, 0.065
    for y, decision_text, verdict in zip(ys, decisions, verdicts):
        box = patches.FancyBboxPatch(
            (dx0, y - dh / 2), dw, dh, boxstyle="round,pad=0.01",
            facecolor="#EAF2F8", edgecolor=PALETTE[0], linewidth=1.2,
        )
        ax.add_patch(box)
        ax.text(dcx, y, decision_text, ha="center", va="center", fontsize=9, linespacing=1.4)
        ax.annotate("", xy=(vx0, y), xytext=(dx0 + dw, y),
                    arrowprops={"arrowstyle": "->", "color": PALETTE[1], "lw": 1.2})
        ax.text((dx0 + dw + vx0) / 2, y + 0.012, "No", ha="center", va="bottom",
                fontsize=8.5, color=PALETTE[1])
        vbox = patches.FancyBboxPatch(
            (vx0, y - vh / 2), vw, vh, boxstyle="round,pad=0.01",
            facecolor="#FDEDEC", edgecolor=PALETTE[1], linewidth=1.0,
        )
        ax.add_patch(vbox)
        ax.text(vx0 + vw / 2, y, verdict, ha="center", va="center", fontsize=9)
        y_next = y - 0.20 if y != ys[-1] else 0.09
        top_next = y_next + (dh / 2 if y != ys[-1] else 0.045)
        ax.annotate("", xy=(dcx, top_next), xytext=(dcx, y - dh / 2),
                    arrowprops={"arrowstyle": "->", "color": PALETTE[2], "lw": 1.2})
        ax.text(dcx + 0.015, (y - dh / 2 + top_next) / 2, "Yes", ha="left", va="center",
                fontsize=8.5, color=PALETTE[2])
    final = patches.FancyBboxPatch(
        (dcx - 0.16, 0.045), 0.32, 0.09, boxstyle="round,pad=0.01",
        facecolor="#E8F8F5", edgecolor=PALETTE[2], linewidth=1.2,
    )
    ax.add_patch(final)
    ax.text(dcx, 0.09, "Robust nonlinear", ha="center", va="center", fontsize=10, weight="bold")
    y_ins = (ys[0] + ys[1]) / 2
    ibox = patches.FancyBboxPatch(
        (vx0, y_ins - vh / 2), vw, vh, boxstyle="round,pad=0.01",
        facecolor="#FEF9E7", edgecolor=PALETTE[4], linewidth=1.0,
    )
    ax.add_patch(ibox)
    ax.text(vx0 + vw / 2, y_ins, "Insufficient data", ha="center", va="center", fontsize=9)
    ax.annotate("", xy=(vx0, y_ins), xytext=(dx0 + dw, ys[1] + dh / 2 - 0.02),
                arrowprops={"arrowstyle": "->", "color": PALETTE[4], "lw": 1.2,
                            "connectionstyle": "arc3,rad=0.3"})
    save_fig(fig, "Figure2")
    paths["fig_decision_tree"] = FIGURES / "Figure2.png"

    sens = pd.read_csv(RESULTS / "verdict_sensitivity_summary.csv")
    heat = sens.pivot(index="alpha", columns="k", values="share_not_robust") * 100
    fig, ax = plt.subplots(figsize=(6.4, 4.2))
    im = ax.imshow(heat.to_numpy(), cmap="coolwarm", vmin=0, vmax=100, aspect="auto")
    ax.set_xticks(range(len(heat.columns)), [str(x) for x in heat.columns])
    ax.set_yticks(range(len(heat.index)), [f"{x:.2f}" for x in heat.index])
    ax.set_xlabel("Cook-distance removal count (k)")
    ax.set_ylabel("Significance level (alpha)")
    ax.set_title("Share not classified as robustly nonlinear")
    for i in range(len(heat.index)):
        for j in range(len(heat.columns)):
            ax.text(j, i, f"{heat.iloc[i, j]:.0f}%", ha="center", va="center", color="black")
    fig.colorbar(im, ax=ax, label="Share not robust (%)")
    fig.tight_layout()
    save_fig(fig, "Figure3")
    paths["fig_sensitivity"] = FIGURES / "Figure3.png"

    desc = pd.read_csv(FUKU_RESULTS / "descriptive_cdr.csv")
    fig, ax = plt.subplots(figsize=(7.2, 4.3))
    labels = {"T1": "T1", "T1w": "T1w", "T2": "T2", "control": "Fukushima controls"}
    for i, (group, frame) in enumerate(desc.groupby("group")):
        ax.plot(frame.year, frame.cdr, label=labels[group], color=PALETTE[i], linewidth=2)
    ax.axvspan(2012, 2019, color=PALETTE[5], alpha=0.12, label="ATT window (2012-2019)")
    ax.axvline(2011, color="0.35", linestyle="--", linewidth=1)
    ax.set_xticks(np.arange(2005, 2023, 2))
    ax.set(xlabel="Year", ylabel="Crude death rate per 1,000", title="Pooled crude death rates")
    ax.legend(frameon=False, ncol=2, loc="upper center", bbox_to_anchor=(0.5, -0.18))
    fig.tight_layout(rect=(0, 0.08, 1, 1))
    save_fig(fig, "Figure4")
    paths["fig_pooled_cdr"] = FIGURES / "Figure4.png"

    event = pd.read_csv(FUKU_RESULTS / "did_event_study.csv")
    frame = event.query(
        "treatment == 'T1' and control_pool == 'fukushima_only' "
        "and weighting == 'unweighted' and age_control == False and outcome == 'ln_cdr'"
    ).sort_values("event_year")
    fig, ax = plt.subplots(figsize=(7.2, 4.3))
    ax.axhline(0, color="0.2", linewidth=0.8)
    ax.axvspan(2012, 2019, color=PALETTE[5], alpha=0.12, label="ATT window (2012-2019)")
    ax.axvline(2011, color="0.35", linestyle="--", linewidth=1)
    ax.text(2011.1, ax.get_ylim()[1] * 0.9, "2011 (excluded from ATT)", fontsize=8, rotation=90, va="top")
    ax.errorbar(frame.event_year, frame.beta, yerr=1.96 * frame.se, fmt="o-", color=PALETTE[0], capsize=3)
    ax.set(xlabel="Event year", ylabel="Log CDR coefficient", title="T1 event study: Fukushima controls")
    fig.tight_layout()
    save_fig(fig, "Figure5")
    paths["fig_event_study"] = FIGURES / "Figure5.png"

    acute = pd.read_csv(FUKU_RESULTS / "acute_2011_decomposition.csv")
    row = acute.query("group == 'T1' and muni_name == 'pooled'").iloc[0]
    fig, ax = plt.subplots(figsize=(6.4, 4.2))
    expected = float(row.expected_2011)
    direct = float(row.direct_deaths_2011)
    residual = float(row.residual_2011)
    ax.bar(["Expected 2011"], [expected], color=PALETTE[0], label="Expected")
    ax.bar(["Observed 2011"], [expected], color=PALETTE[0], label="Expected")
    ax.bar(["Observed 2011"], [direct], bottom=[expected], color=PALETTE[1], label="Direct disaster deaths")
    ax.bar(["Observed 2011"], [residual], bottom=[expected + direct], color=PALETTE[2], label="Residual excess")
    for x, y, label in [(0, expected, expected), (1, expected, expected), (1, direct, direct), (1, residual, residual)]:
        bottom = 0 if x == 0 else (0 if label == expected else expected if label == direct else expected + direct)
        ax.text(x, bottom + y / 2, f"{label:,.0f}", ha="center", va="center", fontsize=8)
    ax.set_ylabel("Deaths")
    ax.set_title("T1 acute-phase mortality decomposition, 2011")
    fig.tight_layout()
    save_fig(fig, "Figure6")
    paths["fig_acute"] = FIGURES / "Figure6.png"

    ks_panel = pd.read_csv(KANSAS_DATA / "state_panel.csv")
    cpi_2012 = float(ks_panel.loc[ks_panel["year"] == 2012, "cpi"].mean())
    ks_panel["ln_rgdp_pc"] = np.log(ks_panel["rgdp"] * 1e6 / (ks_panel["pop"] * 1e3))
    ks_panel["ln_emp"] = np.log(ks_panel["nonfarm_emp"])
    ks_panel["real_iit_pc"] = ks_panel["tax_indiv_inc"] / ks_panel["pop"] / (ks_panel["cpi"] / cpi_2012)
    ks_panel["ln_real_iit_pc"] = np.log(ks_panel["real_iit_pc"].where(ks_panel["real_iit_pc"] > 0))
    ks_panel["real_tax_pc"] = ks_panel["tax_total"] / ks_panel["pop"] / (ks_panel["cpi"] / cpi_2012)
    ks_panel["ln_real_tax_pc"] = np.log(ks_panel["real_tax_pc"].where(ks_panel["real_tax_pc"] > 0))
    iit_w = pd.read_csv(KANSAS_RESULTS / "sc_weights_ln_real_iit_pc.csv")
    gdp_w = pd.read_csv(KANSAS_RESULTS / "sc_weights_ln_rgdp_pc.csv")
    fig, axes = plt.subplots(1, 2, figsize=(9.2, 3.8), sharex=True)
    for ax, outcome, weights, ylabel in [
        (axes[0], "ln_real_iit_pc", iit_w, "Log real IIT revenue per capita"),
        (axes[1], "ln_rgdp_pc", gdp_w, "Log real GDP per capita"),
    ]:
        weights = weights.rename(columns={weights.columns[0]: "donor"})
        donor_weights = dict(zip(weights.donor, weights.weight))
        sub = ks_panel[ks_panel.state.isin(["KS"] + list(donor_weights))].pivot(index="year", columns="state", values=outcome)
        synthetic = sum(sub[d] * donor_weights[d] for d in donor_weights if d in sub)
        ax.plot(sub.index, sub["KS"], color=PALETTE[1], label="Kansas", linewidth=2)
        ax.plot(sub.index, synthetic, color=PALETTE[0], label="Synthetic Kansas", linewidth=2)
        ax.axvspan(2013, 2017, color=PALETTE[5], alpha=0.12)
        ax.axvline(2012, color="0.35", linestyle="--", linewidth=1, label="2012 enactment")
        ax.axvline(2017, color="0.35", linestyle=":", linewidth=1, label="2017 reversal")
        ax.set_xlabel("Year")
        ax.set_ylabel(ylabel)
        ax.legend(frameon=False, fontsize=8)
        ax.set_title("(a) Real individual income-tax revenue per capita" if outcome == "ln_real_iit_pc" else "(b) Real GDP per capita")
    fig.suptitle("Kansas versus synthetic Kansas")
    fig.tight_layout()
    save_fig(fig, "Figure7")
    paths["fig_kansas"] = FIGURES / "Figure7.png"

    fig, ax = plt.subplots(figsize=(7.2, 8.3))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    if JOURNAL == "jcpa":
        comparative_stages = [
            (
                "Empirical relationship",
                "Dose-response  |  Tax rate-revenue",
                "#EAF2F8",
                PALETTE[0],
            ),
            (
                "Decision-relevant shape",
                "Low-dose slope  |  Local slope and position",
                "#EBF5FB",
                PALETTE[0],
            ),
            (
                "Evidence in policy range",
                "Extrapolation below range  |  Untested rate shift",
                "#E8F8F5",
                PALETTE[2],
            ),
            (
                "Decision sensitivity",
                "Evacuation scope  |  Revenue outlook",
                "#FEF9E7",
                PALETTE[4],
            ),
            (
                "Consequence asymmetry",
                "Exposure vs displacement harms  |  Fiscal vs growth costs",
                "#FDEDEC",
                PALETTE[1],
            ),
            (
                "Reversibility and learning",
                "Staged return, zoning  |  Triggers, sunset, monitoring",
                "#F4ECF7",
                PALETTE[3],
            ),
        ]
        x0, width, height = 0.08, 0.84, 0.105
        ys = [0.86, 0.715, 0.57, 0.425, 0.28, 0.135]
        for index, ((title, subtitle, face, edge), y) in enumerate(zip(comparative_stages, ys)):
            box = patches.FancyBboxPatch(
                (x0, y - height / 2),
                width,
                height,
                boxstyle="round,pad=0.012",
                facecolor=face,
                edgecolor=edge,
                linewidth=1.4,
            )
            ax.add_patch(box)
            ax.text(
                x0 + width / 2,
                y + 0.017,
                title,
                ha="center",
                va="center",
                fontsize=10,
                weight="bold",
            )
            ax.text(
                x0 + width / 2,
                y - 0.022,
                subtitle,
                ha="center",
                va="center",
                fontsize=8,
            )
            if index < len(comparative_stages) - 1:
                ax.annotate(
                    "",
                    xy=(0.5, ys[index + 1] + height / 2 + 0.005),
                    xytext=(0.5, y - height / 2 - 0.005),
                    arrowprops={"arrowstyle": "->", "color": "0.35", "lw": 1.2},
                )
        ax.text(
            0.5,
            0.965,
            "Comparative appraisal sequence: Fukushima evacuation | Kansas tax reform",
            ha="center",
            va="center",
            fontsize=9.5,
            style="italic",
            color="0.30",
        )
        fig.tight_layout()
        save_fig(fig, "Figure1")
        paths["fig_policy_pathway"] = FIGURES / "Figure1.png"
    else:
        stages = [
            (
                "Empirical relationship",
                "Observed support, causal evidence, and uncertainty",
                "#EAF2F8",
                PALETTE[0],
            ),
            (
                "Functional-form assumption",
                "Threshold, turning point, slope, plateau, or monotonicity",
                "#EBF5FB",
                PALETTE[0],
            ),
            (
                "Policy-relevant range",
                "Population, period, institution, and exposure values",
                "#E8F8F5",
                PALETTE[2],
            ),
            (
                "Decision dependence",
                "Would a plausible alternative shape change the action?",
                "#FEF9E7",
                PALETTE[4],
            ),
            (
                "Asymmetric consequences",
                "Who bears losses from underaction or overaction?",
                "#FDEDEC",
                PALETTE[1],
            ),
            (
                "Action architecture",
                "Commit, stage, pilot, adapt, or use precaution",
                "#F4ECF7",
                PALETTE[3],
            ),
        ]
        x0, width, height = 0.12, 0.76, 0.105
        ys = [0.86, 0.715, 0.57, 0.425, 0.28, 0.135]
        for index, ((title, subtitle, face, edge), y) in enumerate(zip(stages, ys)):
            box = patches.FancyBboxPatch(
                (x0, y - height / 2),
                width,
                height,
                boxstyle="round,pad=0.012",
                facecolor=face,
                edgecolor=edge,
                linewidth=1.4,
            )
            ax.add_patch(box)
            ax.text(
                x0 + width / 2,
                y + 0.017,
                title,
                ha="center",
                va="center",
                fontsize=10,
                weight="bold",
            )
            ax.text(
                x0 + width / 2,
                y - 0.022,
                subtitle,
                ha="center",
                va="center",
                fontsize=8.5,
            )
            if index < len(stages) - 1:
                ax.annotate(
                    "",
                    xy=(0.5, ys[index + 1] + height / 2 + 0.005),
                    xytext=(0.5, y - height / 2 - 0.005),
                    arrowprops={"arrowstyle": "->", "color": "0.35", "lw": 1.2},
                )
        ax.text(
            0.965,
            0.50,
            "Transportability and reversibility condition every step",
            rotation=90,
            ha="center",
            va="center",
            fontsize=8.5,
            color="0.35",
        )
        ax.plot([0.925, 0.925], [0.08, 0.92], color="0.65", linewidth=1.0)
        fig.tight_layout()
        save_fig(fig, "Figure1")
        paths["fig_policy_pathway"] = FIGURES / "Figure1.png"

    did = pd.read_csv(QROOT / "kansas" / "results" / "did_event_study.csv")
    neighbors = did.query("pool == 'neighbors'").copy()
    fig, axes = plt.subplots(2, 2, figsize=(8.2, 6.2))
    outcome_labels = {
        "ln_rgdp_pc": "Real GDP per capita",
        "ln_emp": "Nonfarm employment",
        "ln_real_tax_pc": "Real total tax revenue per capita",
        "ln_real_iit_pc": "Real individual income-tax revenue per capita",
    }
    for ax, outcome in zip(axes.flat, ["ln_rgdp_pc", "ln_emp", "ln_real_tax_pc", "ln_real_iit_pc"]):
        sub = neighbors[neighbors.outcome == outcome].sort_values("event_year")
        ax.axhline(0, color="0.2", linewidth=0.7)
        ax.axvspan(2013, 2017, color=PALETTE[5], alpha=0.12)
        ax.axvline(2012, color="0.35", linestyle="--", linewidth=0.8)
        ax.errorbar(sub.event_year, sub.beta, yerr=1.96 * sub.se, fmt="o-", color=PALETTE[2], capsize=2)
        ax.set_title(outcome_labels[outcome])
        ax.set_xticks(np.arange(1997, 2020, 4))
        ax.set_xlabel("Year")
        ax.set_ylabel("Coefficient")
    fig.suptitle("Kansas neighboring-state DiD event studies")
    fig.tight_layout()
    save_fig(fig, "FigureA1")
    paths["fig_a1"] = FIGURES / "FigureA1.png"
    return paths


def table_data(numbers: dict[str, Any]) -> dict[str, pd.DataFrame]:
    summary = pd.read_csv(RESULTS / "summary_table.csv")
    verdicts = [
        "ROBUST_NONLINEAR",
        "NOT_SIGNIFICANT",
        "OUTLIER_DEPENDENT",
        "OVERFITTING",
        "INSUFFICIENT_DATA",
    ]
    labels = {
        "ROBUST_NONLINEAR": "Robust nonlinear",
        "NOT_SIGNIFICANT": "Not significant",
        "OUTLIER_DEPENDENT": "Outlier dependent",
        "OVERFITTING": "Overfitting",
        "INSUFFICIENT_DATA": "Insufficient data",
    }
    audit = summary.groupby(["Category", "Verdict"]).size().unstack(fill_value=0)
    audit = audit.reindex(columns=verdicts, fill_value=0).reset_index()
    audit.columns = ["Category"] + [labels[v] for v in verdicts]
    audit["Total"] = audit.iloc[:, 1:].sum(axis=1)
    total = {"Category": "Total", **{c: int(audit[c].sum()) for c in audit.columns if c != "Category"}}
    audit = pd.concat([audit, pd.DataFrame([total])], ignore_index=True)

    perm = pd.read_csv(FUKU_RESULTS / "did_permutation_inference.csv")
    panel = pd.read_csv(FUKU_DATA / "municipal_panel.csv", dtype={"muni_code": str})
    sample = json.loads((FUKU_RESULTS / "sample_summary.json").read_text())
    treatment_codes = {
        "T1": set(str(x).zfill(5) for x in sample["t1_codes"]),
        "T1w": {"07546", "07545", "07543", "07564"},
        "T2": {"07541", "07209"},
    }
    observed_post = {
        treatment: panel.loc[
            panel["muni_code"].astype(str).str.zfill(5).isin(codes)
            & panel["year"].between(2012, 2019),
            "deaths",
        ].sum()
        for treatment, codes in treatment_codes.items()
    }
    control_labels = {"all_prefectures": "All prefectures", "fukushima_only": "Fukushima only"}
    weighting_labels = {"weighted": "Weighted", "unweighted": "Unweighted"}
    rows = []
    for _, r in perm.query("treatment == 'T1'").iterrows():
        implied = observed_post["T1"] * (1 - math.exp(-float(r.att)))
        rows.append({
            "Group": "T1", "Controls": control_labels[r.control_pool],
            "Weighting": weighting_labels[r.weighting],
            "Age control": "Yes" if r.age_control else "No", "ATT (log points)": f"{r.att:.3f}",
            "Permutation p": ptext(r.perm_p_att), "Pre-trend p": ptext(r.perm_p_pretrend),
            "Implied deaths 2012-2019": comma(implied),
        })
    for treatment in ["T1w", "T2"]:
        for pool in ["all_prefectures", "fukushima_only"]:
            r = perm.query(
                "treatment == @treatment and control_pool == @pool "
                "and weighting == 'unweighted' and age_control == False"
            ).iloc[0]
            implied = observed_post[treatment] * (1 - math.exp(-float(r.att)))
            rows.append({
                "Group": treatment, "Controls": control_labels[pool], "Weighting": "Unweighted",
                "Age control": "No", "ATT (log points)": f"{r.att:.3f}",
                "Permutation p": ptext(r.perm_p_att), "Pre-trend p": ptext(r.perm_p_pretrend),
                "Implied deaths 2012-2019": comma(implied),
            })
    fuku_table = pd.DataFrame(rows)
    lnt = pd.read_csv(FUKU_RESULTS / "lnt_comparison.csv")
    acute = pd.read_csv(FUKU_RESULTS / "acute_2011_decomposition.csv").query("group == 'T1' and muni_name == 'pooled'").iloc[0]
    preferred_t1 = perm.query(
        "treatment == 'T1' and control_pool == 'fukushima_only' "
        "and weighting == 'unweighted' and age_control == False"
    ).iloc[0]
    signed_implied = observed_post["T1"] * (1 - math.exp(-float(preferred_t1.att)))
    who_values = {
        key: lnt.loc[lnt["quantity"].str.contains(pattern, regex=False), "value"].iloc[0]
        for key, pattern in {
            "leukaemia": "WHO leukemia",
            "breast": "WHO breast",
            "solid": "WHO all-solid",
            "thyroid": "WHO thyroid",
        }.items()
    }
    who_conclusion = lnt.loc[lnt["quantity"] == "WHO conclusion outside Japan", "value"].iloc[0]
    unscear = lnt.loc[lnt["quantity"].str.startswith("UNSCEAR collective"), "value"].tolist()
    unscear_conclusion = lnt.loc[lnt["quantity"] == "UNSCEAR conclusion on discernibility", "verbatim_quote"].iloc[0]
    thj_source = lnt.loc[lnt["quantity"].str.startswith("ten Hoeve"), "value"].iloc[0]
    thj_range_match = re.search(r"([\d,]+)\s*[\u2013-]\s*([\d,]+)", str(thj_source))
    thj_best_match = re.search(r"best estimate\s+([\d,]+)", str(thj_source))
    if not thj_range_match or not thj_best_match:
        raise ValueError("Could not parse ten Hoeve and Jacobson estimate from lnt_comparison.csv")
    thj_value = int(thj_best_match.group(1).replace(",", ""))
    thj_range = (
        f"{int(thj_range_match.group(1).replace(',', '')):,}-"
        f"{int(thj_range_match.group(2).replace(',', '')):,}"
    )
    thj_avoided = lnt.loc[
        lnt["quantity"].str.startswith("ten Hoeve and Jacobson evacuation-avoided"),
    ].iloc[0]
    avoided_value = str(thj_avoided["value"])
    avoided_best = re.search(r"best estimate ([^)]+)", avoided_value).group(1)
    lnt_table = pd.DataFrame([
        {"Quantity": "T1 observed deaths 2011", "Value": comma(acute.deaths_2011), "Scope": "T1 municipalities", "Type": "Measured", "Source": "This study"},
        {"Quantity": "T1 expected deaths 2011", "Value": comma(acute.expected_2011), "Scope": "T1 municipalities", "Type": "Measured", "Source": "This study"},
        {"Quantity": "T1 excess deaths 2011", "Value": comma(acute.excess_2011), "Scope": "T1 municipalities", "Type": "Measured", "Source": "This study"},
        {"Quantity": "Of which direct disaster deaths", "Value": comma(acute.direct_deaths_2011), "Scope": "T1 municipalities", "Type": "Measured", "Source": f"Fukushima Prefecture, {numbers['fdr_year']}"},
        {"Quantity": "T1 residual excess deaths 2011", "Value": comma(acute.residual_2011), "Scope": "T1 municipalities", "Type": "Measured", "Source": "This study"},
        {"Quantity": "T1 implied difference 2012-2019 (recorded)", "Value": comma(signed_implied), "Scope": "T1 municipalities", "Type": "Measured", "Source": "This study"},
        {"Quantity": "Certified disaster-related deaths, Fukushima Prefecture", "Value": f"{numbers['ra_related_deaths']:,}", "Scope": "Fukushima prefecture", "Type": "Measured", "Source": f"Reconstruction Agency, {numbers['ra_year']}"},
        {"Quantity": "Projected worldwide cancer deaths from the release", "Value": f"{thj_value:,} ({thj_range})", "Scope": "Worldwide", "Type": "Projected (LNT-based)", "Source": "ten Hoeve & Jacobson, 2012"},
        {"Quantity": "Evacuation-avoided radiological mortalities", "Value": f"{avoided_best} ({avoided_value.split('(')[0].strip()})", "Scope": "Worldwide", "Type": "Projected (LNT-based)", "Source": "ten Hoeve & Jacobson, 2012"},
        {"Quantity": "Lifetime-risk relative increases in the most affected areas", "Value": f"leukaemia {who_values['leukaemia']}; breast {who_values['breast']}; all solid cancers {who_values['solid']}; thyroid {who_values['thyroid']}", "Scope": "Most affected areas", "Type": "Projected (LNT-based)", "Source": "WHO, 2013"},
        {"Quantity": "WHO conclusion outside Japan", "Value": who_conclusion, "Scope": "Outside Japan", "Type": "Projected (LNT-based)", "Source": "WHO, 2013"},
        {"Quantity": "Collective effective dose", "Value": f"{' / '.join(str(int(float(v))) for v in unscear)} thousand man-Sv (first year / ten years / lifetime)", "Scope": "Japanese population", "Type": "Projected (LNT-based)", "Source": "UNSCEAR, 2014"},
        {"Quantity": "UNSCEAR conclusion", "Value": "no discernible increase expected", "Scope": "Exposed population", "Type": "Projected (LNT-based)", "Source": "UNSCEAR, 2014"},
    ])
    for frame in [audit, fuku_table, lnt_table]:
        for column in frame.columns:
            frame[column] = frame[column].map(ascii_clean)

    ks = pd.read_csv(KANSAS_RESULTS / "sc_summary.csv")
    placebo = pd.read_csv(KANSAS_RESULTS / "placebo_in_time.csv").set_index("outcome")
    loo = pd.read_csv(KANSAS_RESULTS / "loo_robustness.csv")
    krows = []
    for _, r in ks.iterrows():
        subset = loo[loo.outcome == r.outcome]
        krows.append({
            "Outcome": {
                "ln_rgdp_pc": "Real GDP per capita",
                "ln_emp": "Nonfarm employment",
                "ln_real_tax_pc": "Real total tax revenue per capita",
                "ln_real_iit_pc": "Real individual income-tax revenue per capita",
            }[r.outcome],
            "Pool": "Main" if r.pool == "main" else "All states",
            "Mean gap": f"{r.post_mean_gap:.3f}",
            "Gap (%)": f"{100 * (math.exp(r.post_mean_gap) - 1):+.1f}%",
            "Permutation p": f"{r.perm_p:.2f}", "Pre-RMSPE": f"{r.rmspe_pre:.3f}",
            "donors": int(r.n_donors), "in-time placebo": f"{placebo.loc[r.outcome, 'post_mean_gap']:.3f}",
            "LOO range": f"{subset.post_mean_gap.min():.3f} to {subset.post_mean_gap.max():.3f}" if r.pool == "main" and not subset.empty else "-",
        })
    kansas_table = pd.DataFrame(krows)
    exclusions = pd.read_csv(FUKU_RESULTS / "dropped_units.csv")
    a1 = exclusions.rename(
        columns={"pref_name": "Prefecture", "muni_name_romaji": "Municipality", "reason": "Reason", "years_missing": "Years missing"}
    )[["Prefecture", "Municipality", "Reason", "Years missing"]].copy()
    a2 = pd.read_csv(FUKU_RESULTS / "sc_unit_summary.csv").copy()
    a2["Municipality"] = a2.muni_code.astype(str).str.zfill(5).map(ROMAJI_BY_CODE).fillna("Municipality")
    a2 = a2.rename(columns={"n_donors": "Donors", "rmspe_pre": "Pre-RMSPE", "rmspe_post": "Post-RMSPE", "post_mean_gap": "Mean gap 2012-2019", "perm_p": "Permutation p"})
    a2 = a2[["Municipality", "Donors", "Pre-RMSPE", "Post-RMSPE", "Mean gap 2012-2019", "Permutation p"]].copy()
    for frame in [kansas_table, a1, a2]:
        for column in frame.columns:
            frame[column] = frame[column].map(ascii_clean)
    if JOURNAL == "jcpa":
        synthesis = pd.DataFrame([
            {"Comparative dimension": "Policy domain", "Fukushima evacuation": "Radiation protection; emergency public health", "Kansas tax reform": "Fiscal and economic policy", "Comparative lesson": "The mechanism is domain-general: a named curve can carry a decision in any sector."},
            {"Comparative dimension": "Policy objective", "Fukushima evacuation": "Reduce population radiation exposure", "Kansas tax reform": "Raise growth while recovering revenue", "Comparative lesson": "Both objectives required a specific shape, not just an association."},
            {"Comparative dimension": "Canonical relationship", "Fukushima evacuation": "Low-dose radiation dose-response", "Kansas tax reform": "Tax rate-revenue and growth response", "Comparative lesson": "Named curves compress multi-step arguments into portable pictures."},
            {"Comparative dimension": "Decision-relevant feature", "Fukushima evacuation": "Low-dose slope below observed doses", "Kansas tax reform": "Local position, slope, and dynamic response", "Comparative lesson": "Each decision needed one feature of the curve, not the curve as a whole."},
            {"Comparative dimension": "Type of uncertainty", "Fukushima evacuation": "Extrapolation below the observed range", "Kansas tax reform": "Local-position and response-geometry uncertainty", "Comparative lesson": "Different uncertainty locations produce different appraisal priorities."},
            {"Comparative dimension": "Extrapolation burden", "Fukushima evacuation": "High; few observations below 0.1 Gy", "Kansas tax reform": "Moderate; within-state history but untested rate cut", "Comparative lesson": "Extrapolation burden predicts how far evidence can carry a decision."},
            {"Comparative dimension": "Transportability", "Fukushima evacuation": "Atomic-bomb and occupational evidence to municipal exposure", "Kansas tax reform": "Cross-state and theoretical evidence to one state's response", "Comparative lesson": "Transport gaps differ in kind; decompose them by population, period, and institution."},
            {"Comparative dimension": "Decision sensitivity", "Fukushima evacuation": "Evacuation scope and duration change under plausible dose-response shapes", "Kansas tax reform": "Revenue outlook changes under plausible local slopes and growth paths", "Comparative lesson": "In both, plausible alternative shapes imply materially different policies."},
            {"Comparative dimension": "Cost of wrong action", "Fukushima evacuation": "Displacement, mortality risk, care disruption", "Kansas tax reform": "Revenue loss and service compression", "Comparative lesson": "Error costs fall on identifiable groups, not on the model."},
            {"Comparative dimension": "Cost of inaction", "Fukushima evacuation": "Residual radiation exposure", "Kansas tax reform": "Foregone growth stimulus claim", "Comparative lesson": "Asymmetry must be weighed in both directions, not one."},
            {"Comparative dimension": "Consequence asymmetry", "Fukushima evacuation": "Severe on both sides; competing harms", "Kansas tax reform": "Primarily fiscal; distribution across budgets", "Comparative lesson": "Higher asymmetry raises the decision value of shape certainty."},
            {"Comparative dimension": "Reversibility", "Fukushima evacuation": "Limited; displacement and health effects partly irreversible", "Kansas tax reform": "Comparatively greater; statutes revised in 2017", "Comparative lesson": "Reversible settings can substitute adaptive design for ex ante certainty."},
            {"Comparative dimension": "Time to observe outcomes", "Fukushima evacuation": "Years; health endpoints partly latent", "Kansas tax reform": "Years; fiscal outcomes observable quickly", "Comparative lesson": "Observability on policy timescales enables learning architectures."},
            {"Comparative dimension": "Learning opportunity", "Fukushima evacuation": "Zoning, staged return, exemptions possible but hard mid-crisis", "Kansas tax reform": "Staged cuts, triggers, sunset clauses readily available", "Comparative lesson": "Where learning is feasible, adaptive implementation hedges shape uncertainty."},
            {"Comparative dimension": "Preferred decision architecture", "Fukushima evacuation": "Staged, revisable protection with asymmetric-harm safeguards", "Kansas tax reform": "Staged implementation with triggers and monitoring", "Comparative lesson": "Architecture should match the joint uncertainty-consequence-reversibility configuration."},
        ])
    else:
        synthesis = pd.DataFrame([
            {"Dimension": "Policy domain", "Fukushima": "Health and radiation protection", "Kansas": "Fiscal policy"},
            {"Dimension": "Canonical relationship", "Fukushima": "Dose-response", "Kansas": "Tax rate-revenue"},
            {"Dimension": "Decision-relevant feature", "Fukushima": "Low-dose slope and extrapolation", "Kansas": "Local slope and turning point"},
            {"Dimension": "Central uncertainty", "Fukushima": "Extrapolation below the observed range", "Kansas": "Unknown position and behavioral response"},
            {"Dimension": "Principal downside of error", "Fukushima": "Health and displacement harms", "Kansas": "Revenue and growth effects"},
            {"Dimension": "Reversibility", "Fukushima": "Relatively limited", "Kansas": "Comparatively greater"},
            {"Dimension": "Functional-form dependence", "Fukushima": "High if action changes under alternative plausible shapes", "Kansas": "High if expected fiscal response depends on local curve geometry"},
        ])
    for column in synthesis.columns:
        synthesis[column] = synthesis[column].map(ascii_clean)
    return {"tab_audit": audit, "tab_fuku_att": fuku_table, "tab_lnt": lnt_table, "tab_kansas": kansas_table, "tab_synthesis": synthesis, "tab_a1": a1, "tab_a2": a2}


def add_table(doc: Document, frame: pd.DataFrame) -> None:
    table = doc.add_table(rows=1, cols=len(frame.columns))
    table.style = "Table Grid"
    for cell, name in zip(table.rows[0].cells, frame.columns):
        cell.text = str(name)
    for _, row in frame.iterrows():
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = str(value)


def omml_run(symbol: str, subscript: str | None = None) -> OxmlElement:
    omath = OxmlElement("m:oMath")
    if subscript:
        ssub = OxmlElement("m:sSub")
        base = OxmlElement("m:e")
        r = OxmlElement("m:r")
        t = OxmlElement("m:t")
        t.text = symbol
        r.append(t)
        base.append(r)
        sub = OxmlElement("m:sub")
        r2 = OxmlElement("m:r")
        t2 = OxmlElement("m:t")
        t2.text = subscript
        r2.append(t2)
        sub.append(r2)
        ssub.extend([base, sub])
        omath.append(ssub)
    else:
        r = OxmlElement("m:r")
        t = OxmlElement("m:t")
        t.text = symbol
        r.append(t)
        omath.append(r)
    return omath


def add_inline_math(paragraph, token: str) -> None:
    mapping = {"a_i": ("alpha", "i"), "g_t": ("gamma", "t"), "b": ("beta", None), "e_it": ("epsilon", "it")}
    symbol, sub = mapping[token]
    paragraph._p.append(omml_run(symbol, sub))


def add_text_with_math(paragraph, text: str) -> None:
    pattern = re.compile(r"(a_i|g_t|e_it|(?<![A-Za-z])b(?![A-Za-z]))")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor:match.start()])
        add_inline_math(paragraph, match.group(0))
        cursor = match.end()
    paragraph.add_run(text[cursor:])


def add_block_equation(doc: Document, expression: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    omath_para = OxmlElement("m:oMathPara")
    omath = OxmlElement("m:oMath")
    for token in ["ln(CDR", "_it) = ", "alpha_i", " + ", "gamma_t", " + ", "beta", " D_it + ", "epsilon_it"]:
        if token == "alpha_i":
            omath.append(omml_run("alpha", "i"))
        elif token == "gamma_t":
            omath.append(omml_run("gamma", "t"))
        elif token == "beta":
            omath.append(omml_run("beta"))
        elif token == "epsilon_it":
            omath.append(omml_run("epsilon", "it"))
        else:
            r = OxmlElement("m:r")
            t = OxmlElement("m:t")
            t.text = token
            r.append(t)
            omath.append(r)
    omath_para.append(omath)
    p._p.append(omath_para)


def add_heading_or_para(doc: Document, kind: str, text: str, numbers: dict[str, Any]) -> None:
    text = format_text(text, numbers)
    if kind == "h1":
        doc.add_heading(text, level=1)
    elif kind == "h2":
        doc.add_heading(text, level=2)
    elif kind == "p":
        p = doc.add_paragraph()
        add_text_with_math(p, text)
    elif kind == "eq":
        add_block_equation(doc, text)


def build_docx(numbers: dict[str, Any], tables: dict[str, pd.DataFrame], figures: dict[str, Path]) -> None:
    def render(path: Path, inline: bool, separate_tables: bool = False) -> None:
        doc = Document()
        configure_document(doc)
        doc.core_properties.author = ""
        doc.core_properties.last_modified_by = ""
        doc.add_heading(format_text(text.TITLE, numbers), level=0)
        p = doc.add_paragraph()
        p.add_run("Abstract").bold = True
        doc.add_paragraph(format_text(text.ABSTRACT, numbers))
        p = doc.add_paragraph()
        p.add_run("Keywords: ").bold = True
        p.add_run(text.KEYWORDS)
        table_counter = 0
        figure_counter = 0
        deferred_tables: list[tuple[str, pd.DataFrame]] = []
        deferred_figures: list[tuple[str, str]] = []
        for kind, value in text.BODY:
            if kind in {"h1", "h2", "p", "eq"}:
                add_heading_or_para(doc, kind, value or "", numbers)
            elif kind == "refs":
                doc.add_heading("References", level=1)
                for ref in sorted(
                    format_text(r, numbers) for r in text.REFERENCES
                ):
                    doc.add_paragraph(ref)
            elif kind == "fig":
                figure_counter += 1
                if inline:
                    doc.add_picture(str(figures[value]), width=Inches(6.4))
                    cap = doc.add_paragraph(
                        format_text(text.CAPTIONS[value], numbers),
                        style="Caption",
                    )
                    cap.paragraph_format.space_before = Pt(12)
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    deferred_figures.append(
                        (
                            value,
                            format_text(text.CAPTIONS[value], numbers),
                        )
                    )
                    doc.add_paragraph(f"Figure {figure_counter} about here")
            elif kind == "table":
                table_counter += 1
                if inline:
                    cap = doc.add_paragraph(
                        format_text(text.CAPTIONS[value], numbers),
                        style="Caption",
                    )
                    cap.paragraph_format.space_before = Pt(12)
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    add_table(doc, tables[value])
                else:
                    deferred_tables.append((value, tables[value]))
        if not inline:
            doc.add_heading("Figure legends", level=1)
            for _, caption in deferred_figures:
                doc.add_paragraph(caption)
            doc.add_heading("Tables", level=1)
            for key, frame in deferred_tables:
                doc.add_paragraph(
                    format_text(text.CAPTIONS[key], numbers),
                    style="Caption",
                )
                add_table(doc, frame)
        if JOURNAL == "jcpa":
            doc.add_page_break()
            doc.add_heading(
                format_text(text.APPENDIX[0][1], numbers), level=1
            )
            for kind, value in text.APPENDIX[1:]:
                doc.add_paragraph(
                    format_text(text.CAPTIONS[value], numbers),
                    style="Caption",
                )
                if kind == "table":
                    add_table(doc, tables[value])
                elif kind == "fig":
                    doc.add_picture(
                        str(figures[value]), width=Inches(6.4)
                    )
        path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(path)

    render(ROOT / f"manuscript_{JOURNAL}.docx", True)
    render(ROOT / f"manuscript_{JOURNAL}_figures_separate.docx", False)
    numbers["manuscript_word_count"] = document_word_count(
        ROOT / f"manuscript_{JOURNAL}.docx"
    )
    (MANUSCRIPT / "numbers.json").write_text(
        json.dumps(numbers, indent=2), encoding="utf-8"
    )

    appendix = Document()
    configure_document(appendix)
    appendix.core_properties.author = ""
    appendix.core_properties.last_modified_by = ""
    appendix.add_heading(
        format_text(text.APPENDIX[0][1], numbers), level=0
    )
    for kind, value in text.APPENDIX[1:]:
        appendix.add_paragraph(
            format_text(text.CAPTIONS[value], numbers),
            style="Caption",
        )
        if kind == "table":
            add_table(appendix, tables[value])
        elif kind == "fig":
            appendix.add_picture(str(figures[value]), width=Inches(6.4))
    appendix.save(ROOT / f"appendix_{JOURNAL}.docx")

    title = Document()
    configure_document(title)
    meta = JOURNAL_META[JOURNAL]
    title.add_heading(text.TITLE, level=0)
    title.add_paragraph("Tatsuki Onishi")
    title.add_paragraph("Data Science and AI Innovation Research Promotion Center, Shiga University")
    title.add_paragraph("Corresponding author: bougtoir@gmail.com")
    title.add_paragraph(f"Journal: {meta['name']}")
    title.add_paragraph(f"Article type: {meta['article_type']}")
    title.add_paragraph(f"Manuscript word count (all inclusive): {numbers['manuscript_word_count']:,}")
    title.add_paragraph(f"Data and code: {numbers['repo_url']}")
    title.add_paragraph("Funding: No external funding.")
    title.add_paragraph("Competing interests: The author declares no competing interests.")
    title.add_paragraph(
        f"{text.AI_ASSISTANCE_HEADING}. "
        f"{text.AI_ASSISTANCE_TEXT}"
    )
    title.save(ROOT / f"title_page_{JOURNAL}.docx")

    statements = Document()
    configure_document(statements)
    statements.add_heading(f"{meta['name']} submission statements", level=0)
    statements.add_heading("Author contribution", level=1)
    statements.add_paragraph(
        "Tatsuki Onishi conceived the study, specified the methodology, curated and "
        "verified the data, interpreted the results, and approved the final manuscript."
    )
    statements.add_heading("Funding", level=1)
    statements.add_paragraph("The author received no external funding for this work.")
    statements.add_heading("Competing interests", level=1)
    statements.add_paragraph("The author declares no competing interests.")
    statements.add_heading("Ethics approval", level=1)
    statements.add_paragraph(
        "Not applicable. The study uses publicly available aggregate and administrative data."
    )
    statements.add_heading("Data availability", level=1)
    statements.add_paragraph(
        "Archived public source files, acquisition metadata, processed datasets, analysis "
        f"code, and document-generation code are available at {numbers['repo_url']}. A "
        "double-anonymous reviewer archive is included with the submission."
    )
    statements.add_heading(text.AI_ASSISTANCE_HEADING, level=1)
    statements.add_paragraph(text.AI_ASSISTANCE_TEXT)
    statements.save(ROOT / f"submission_statements_{JOURNAL}.docx")

    tables_doc = Document()
    configure_document(tables_doc)
    tables_doc.add_heading(f"{meta['name']} tables", level=0)
    for key in TABLE_KEYS:
        if key in tables and key in text.CAPTIONS:
            tables_doc.add_paragraph(
                format_text(text.CAPTIONS[key], numbers),
                style="Caption",
            )
            add_table(tables_doc, tables[key])
    tables_doc.save(ROOT / f"tables_{JOURNAL}.docx")


def build_pptx(numbers: dict[str, Any], tables: dict[str, pd.DataFrame], figures: dict[str, Path]) -> None:
    prs = Presentation()
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)
    blank = prs.slide_layouts[6]
    keys = [
        "fig_policy_pathway",
        "fig_decision_tree",
        "fig_sensitivity",
        "fig_pooled_cdr",
        "fig_event_study",
        "fig_acute",
        "fig_kansas",
        "fig_a1",
    ]
    for key in keys:
        slide = prs.slides.add_slide(blank)
        title = slide.shapes.add_textbox(PptInches(0.4), PptInches(0.2), PptInches(12.5), PptInches(0.4))
        title.text_frame.text = format_text(
            text.CAPTIONS[key], numbers
        ).split(". ", 1)[0]
        title.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        title.text_frame.paragraphs[0].runs[0].font.size = PptPt(16)
        with Image.open(figures[key]) as image:
            image_ratio = image.width / image.height
        box_width, box_height = 12.0, 5.45
        if image_ratio >= box_width / box_height:
            image_width = box_width
            image_height = box_width / image_ratio
        else:
            image_height = box_height
            image_width = box_height * image_ratio
        image_left = (13.333 - image_width) / 2
        image_top = 0.65 + (box_height - image_height) / 2
        slide.shapes.add_picture(
            str(figures[key]),
            PptInches(image_left),
            PptInches(image_top),
            width=PptInches(image_width),
            height=PptInches(image_height),
        )
        cap = slide.shapes.add_textbox(PptInches(0.5), PptInches(6.18), PptInches(12.3), PptInches(1.05))
        cap.text_frame.text = format_text(
            text.CAPTIONS[key], numbers
        )
        cap.text_frame.word_wrap = True
        cap.text_frame.paragraphs[0].runs[0].font.size = PptPt(9)
    for key in TABLE_KEYS:
        if key not in tables or key not in text.CAPTIONS:
            continue
        frame = tables[key]
        slide = prs.slides.add_slide(blank)
        title = slide.shapes.add_textbox(PptInches(0.4), PptInches(0.2), PptInches(12.5), PptInches(0.4))
        caption = format_text(text.CAPTIONS[key], numbers)
        title_parts = caption.split(". ", 2)
        title.text_frame.text = ". ".join(title_parts[:2])
        title.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        title.text_frame.paragraphs[0].runs[0].font.size = PptPt(14)
        rows, cols = min(len(frame) + 1, 28), len(frame.columns)
        table = slide.shapes.add_table(rows, cols, PptInches(0.3), PptInches(0.8), PptInches(12.7), PptInches(5.65)).table
        for j, col in enumerate(frame.columns):
            table.cell(0, j).text = str(col)
        for i, (_, row) in enumerate(frame.head(rows - 1).iterrows(), start=1):
            for j, value in enumerate(row):
                table.cell(i, j).text = str(value)
        cap = slide.shapes.add_textbox(PptInches(0.5), PptInches(6.55), PptInches(12.3), PptInches(0.65))
        cap.text_frame.text = caption
        cap.text_frame.word_wrap = True
        cap.text_frame.paragraphs[0].runs[0].font.size = PptPt(8)
    prs.save(ROOT / f"figures_{JOURNAL}.pptx")


COVER_LETTER_BODY = {
    "policy_sciences": [
        (
            "I submit the Analysis article, \"{title}.\" The manuscript "
            "develops functional-form dependence as a decision-facing problem in evidence-based "
            "policymaking: a policy may depend on a threshold, turning point, or response shape "
            "that is not identified over the range where the policy operates."
        ),
        (
            "The article develops a decision-facing appraisal framework, uses a purposive "
            "cross-domain mapping as motivation, and then demonstrates different parts of the "
            "framework through Fukushima evacuation and Kansas tax reform. The framework links "
            "association, shape identification, policy-relevant range, decision dependence, "
            "consequence asymmetry, transportability, and reversibility."
        ),
        (
            "Policy Sciences is the appropriate venue because the contribution is problem-"
            "oriented, connects evidence use to policy design under uncertainty, and derives "
            "applied lessons from multi-method analysis rather than presenting the cases as "
            "stand-alone program evaluations."
        ),
        (
            "The manuscript is original and is not under consideration elsewhere. All data "
            "are from public sources, a double-anonymous reviewer archive is included, "
            "reproducibility materials are available at {repo_url}, and I declare "
            "no competing interests or external funding."
        ),
    ],
    "policy_studies": [
        (
            "I submit the original research article, \"{title},\" for consideration "
            "in Policy Studies. The article addresses a practical weakness in evidence-based "
            "policy analysis: a policy recommendation may depend not merely on whether an "
            "empirical association exists, but on whether the specific threshold, slope, "
            "curvature, turning point, or extrapolated response shape that the decision "
            "requires is actually identified over the policy-relevant range."
        ),
        (
            "Existing appraisal practices can miss this problem because they concentrate on "
            "the credibility of an estimate or the strength of a causal design. A study can "
            "be highly credible for estimating an average contrast while providing little "
            "evidence about the feature of the response function that the policy actually "
            "needs. The article names the resulting condition functional-form dependence and "
            "develops a decision-facing appraisal framework that links the empirical "
            "relationship, the assumed shape, the policy-relevant range, decision "
            "sensitivity, consequence asymmetry, transportability, and reversibility. The "
            "framework is offered as a practical appraisal tool for use before or during "
            "policy design, not as a universal policy theory."
        ),
        (
            "Three connected components support the argument. A purposive cross-domain "
            "mapping applies one transparent nonlinear diagnostic to 50 canonical "
            "relationships and shows that fragile shape identification is not confined to a "
            "single policy field. Two substantively different applications then show how the "
            "dependence reaches real decisions: post-Fukushima evacuation under low-dose "
            "radiation-risk assumptions, where extrapolation below the observed range met "
            "asymmetric and partly irreversible harms, and the Kansas tax reform, where a "
            "theoretically possible revenue response did not identify the state's local "
            "position, slope, or dynamic adjustment. A direct cross-case synthesis explains "
            "how the same decision-level condition arises through different mechanisms."
        ),
        (
            "The design is deliberately cross-domain because the problem is cross-domain. "
            "Radiation protection and fiscal policy share no dataset, institution, or "
            "political context, yet both debates treated a named curve as more identified "
            "than the evidence over the relevant range established. That heterogeneity is "
            "what makes the appraisal problem general and why its management - exposing the "
            "shape assumption, testing decisions under plausible alternatives, and choosing "
            "action architectures that preserve the option to learn - should interest Policy "
            "Studies readers concerned with evidence use, policy appraisal, and policy "
            "design under uncertainty."
        ),
        (
            "The manuscript is interdisciplinary, empirically grounded in public data, and "
            "policy-facing; the framework generalizes beyond the two applications. It is "
            "original and is not under consideration elsewhere. All data are from public "
            "sources, a double-anonymous reviewer archive is included, reproducibility "
            "materials are available at {repo_url}, and I declare no competing interests or "
            "external funding."
        ),
    ],
    "jcpa": [
        (
            "I submit the original research article, \"{title},\" for consideration "
            "in the Journal of Comparative Policy Analysis: Research and Practice. The "
            "article asks a comparative question: how does functional-form dependence - "
            "the condition in which a policy's justification rests on an empirical "
            "relationship shape that is insufficiently identified over the policy-relevant "
            "range - enter policy decisions under substantively different evidentiary, "
            "institutional, and consequence structures, and what can comparison across "
            "those structures teach us about policy appraisal under uncertainty?"
        ),
        (
            "The analysis is built on a deliberate maximum-contrast comparison of two "
            "policies that share no sector, institution, or evidentiary tradition: the "
            "post-Fukushima evacuation under low-dose radiation-risk assumptions and the "
            "Kansas tax reform under a claimed fiscal response curve. The cases were "
            "selected because both expose the same decision-level mechanism - policy "
            "commitment depended on functional-form assumptions whose decision-relevant "
            "features were incompletely identified - while differing sharply in the "
            "location of uncertainty, consequence asymmetry, reversibility, and capacity "
            "for adaptive learning. Comparing them shows what a single-case study cannot: "
            "functional-form dependence alone does not dictate a policy response; its "
            "significance depends on the joint configuration of decision sensitivity, "
            "consequence asymmetry, reversibility, transportability, and learning "
            "capacity."
        ),
        (
            "The article contributes to comparative policy analysis in three ways: it "
            "conceptualizes functional-form dependence as a decision-level manifestation "
            "of model-form uncertainty; it identifies shared and context-dependent "
            "mechanisms through a structured focused comparison organized around common "
            "dimensions and a central comparative table; and it derives comparative "
            "propositions and lesson-drawing that distinguish what travels across policy "
            "domains from what remains context-bound. A purposive cross-domain mapping of "
            "50 canonical relationships supplies background motivation, showing that the "
            "gap between association and decision-relevant shape is not unique to either "
            "case, while the comparison itself carries the article's contribution."
        ),
        (
            "The manuscript is original, is not under consideration elsewhere, and is "
            "prepared for double-anonymous review. All data are from public sources, a "
            "double-anonymous reviewer archive is included, reproducibility materials "
            "are available at {repo_url}, and I declare no competing interests or "
            "external funding."
        ),
    ],
}


def build_cover(numbers: dict[str, Any]) -> None:
    doc = Document()
    configure_document(doc)
    doc.add_heading("Cover Letter", level=0)
    doc.add_paragraph("Dear Editor-in-Chief,")
    for paragraph in COVER_LETTER_BODY[JOURNAL]:
        doc.add_paragraph(
            paragraph.format(title=text.TITLE, repo_url=numbers["repo_url"])
        )
    doc.add_paragraph("Sincerely,\nOnishi Tatsuki")
    doc.save(ROOT / f"cover_letter_{JOURNAL}.docx")


def main() -> None:
    numbers = load_numbers()
    tables = table_data(numbers)
    figures = plot_figures(numbers)
    build_docx(numbers, tables, figures)
    build_pptx(numbers, tables, figures)
    build_cover(numbers)
    print(json.dumps(numbers, indent=2))


if __name__ == "__main__":
    main()
